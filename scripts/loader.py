"""Load XLS/XLSX workbooks into plain grids of cleaned strings.

Cells whose font is Kruti Dev are converted to Unicode here, so the
parsers only ever see Unicode text.
"""
import io
import re
import warnings

import openpyxl
import xlrd

from krutidev import kd2uni

warnings.filterwarnings("ignore")

KRUTI_RE = re.compile(r"kruti|kriti|devlys|dev ?lys", re.I)


def _fmt_num(v):
    if isinstance(v, float):
        if v != v:
            return ""
        if v.is_integer():
            return str(int(v))
        return repr(v)
    return str(v)


def _clean(s):
    s = s.replace("\xa0", " ").replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", s).strip()


def _decrypt(raw):
    import msoffcrypto
    f = msoffcrypto.OfficeFile(io.BytesIO(raw))
    f.load_key(password="VelvetSweatshop")
    out = io.BytesIO()
    f.decrypt(out)
    return out.getvalue()


_KD_SIGNS = re.compile("[a-z][;'\\[\\]|&>?]|>|[a-z]Z\\b|\\b(?:dqekj|flag|flg|izk0|fo0|ikVhZ|Hkkjrh;|ernku|LFky|;ksx|dqy|ek/;fed|izkFkfed|xzke|uxj|ifj\"kn|lHkk|fuokZpu|fo\\|ky;|vH;FkhZ|bafM;u|dq0|eks0|Mk0|izlkn|nsoh|dqekjh|feJk|;kno|xqIrk|oekZ|'kekZ|ikloku|vgen|\\[kku|jktiwr|ikaMs;|Bkdqj|frokjh|nqcs|'kqDyk|ekS;Z|ikB'kkyk|fo|ky;)\\b|¼|½|ñ|[a-z][A-Z][a-z]")


def kd_score(s):
    """How much a plain-ASCII string looks like Kruti Dev encoded Hindi."""
    if not s or re.search("[\u0900-\u097F]", s):
        return 0.0
    letters = [c for c in s if c.isascii() and c.isalpha()]
    if len(letters) < 3 or s.upper() == s:
        return 0.0
    k = sum(1 for c in letters if c == "k") / len(letters)
    return k + 0.35 * len(_KD_SIGNS.findall(s))


_EN_PREFIX = re.compile(r"^((?:candidate'?s?\s*name|name\s*of\s*(?:the\s*)?candidate)\s*[-:.]?\s*)(.*)$", re.I)


def _conv(v):
    m = _EN_PREFIX.match(v)
    if m and m.group(2):
        return _clean(m.group(1) + kd2uni(m.group(2)))
    return _clean(kd2uni(v))


def _score(v):
    m = _EN_PREFIX.match(v)
    return kd_score(m.group(2)) if m and m.group(2) else kd_score(v)


def _fix_unflagged_kruti(rows, font_kruti=0):
    """Convert Kruti Dev text stored under a non-Kruti font.

    A sheet is treated as Kruti-encoded when a good share of its longer
    lowercase text cells score high; then only high-scoring cells convert.
    """
    cells = [(r, c) for r, row in enumerate(rows) for c, v in enumerate(row)
             if v and sum(ch.isalpha() and ch.isascii() for ch in v) >= 4 and v.upper() != v]
    if not cells:
        return 0
    hits = [(r, c) for r, c in cells if _score(rows[r][c]) >= 0.2]
    sheet_is_kruti = len(hits) >= max(3, 0.25 * len(cells)) or font_kruti >= 10
    # a cell that scores very high is converted even in a mostly-English sheet
    cut = 0.21 if sheet_is_kruti else 0.5
    n = 0
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if v and _score(v) > cut:
                row[c] = _conv(v)
                n += 1
    return n


class Sheet:
    def __init__(self, name, rows, kd_cells):
        self.name = name
        self.rows = rows          # list[list[str]]
        self.kd_cells = kd_cells  # number of cells converted from Kruti Dev
        self.kd_heuristic = _fix_unflagged_kruti(rows, kd_cells)
        self.ncols = max((len(r) for r in rows), default=0)
        for r in rows:
            r.extend([""] * (self.ncols - len(r)))

    @property
    def nrows(self):
        return len(self.rows)


def _load_xls(raw):
    try:
        wb = xlrd.open_workbook(file_contents=raw, formatting_info=True)
        fmt = True
    except Exception:
        wb = xlrd.open_workbook(file_contents=raw)
        fmt = False
    sheets = []
    for sh in wb.sheets():
        rows, kd = [], 0
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_fmt_num(v))
                    continue
                if isinstance(v, str):
                    if fmt and v.strip():
                        try:
                            font = wb.font_list[wb.xf_list[sh.cell_xf_index(r, c)].font_index].name
                        except Exception:
                            font = ""
                        if KRUTI_RE.search(font or ""):
                            v = kd2uni(v)
                            kd += 1
                    row.append(_clean(v))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append("")
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append("")
                else:
                    row.append(_fmt_num(v))
            rows.append(row)
        sheets.append(Sheet(sh.name, rows, kd))
    return sheets


def _load_xlsx(raw):
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows, kd = [], 0
        for row in ws.iter_rows():
            out = []
            for c in row:
                v = c.value
                if v is None:
                    out.append("")
                elif isinstance(v, str):
                    if v.strip() and c.font is not None and KRUTI_RE.search(c.font.name or ""):
                        v = kd2uni(v)
                        kd += 1
                    out.append(_clean(v))
                elif isinstance(v, (int, float)):
                    out.append(_fmt_num(float(v)) if isinstance(v, float) else str(v))
                else:
                    out.append(_clean(str(v)))
            rows.append(out)
        # trim trailing empty rows
        while rows and not any(rows[-1]):
            rows.pop()
        sheets.append(Sheet(ws.title, rows, kd))
    return sheets


def load_book(path):
    raw = open(path, "rb").read()
    notes = []
    if raw.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            return _load_xls(raw), notes
        except Exception as e:
            if "encrypt" in str(e).lower():
                raw = _decrypt(raw)
                notes.append("decrypted (default Excel password)")
                if raw.startswith(b"PK"):
                    return _load_xlsx(raw), notes
                return _load_xls(raw), notes
            raise
    if raw.startswith(b"PK"):
        if path.lower().endswith(".xls"):
            notes.append("file is XLSX despite .xls extension")
        return _load_xlsx(raw), notes
    raise ValueError("unknown file type")
